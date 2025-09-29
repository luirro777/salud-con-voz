"""
resources.py
Genera workbooks .xlsx con una hoja por cuestionario (nombre = codigo).
Cada hoja: 2 columnas (Label / Valor), A = verbose_name cuando aplique.
Sanea datetimes con tzinfo para evitar errores de openpyxl.
"""
import re
from io import BytesIO
from decimal import Decimal
from datetime import datetime, date, time, timezone

from openpyxl import Workbook
from openpyxl.styles import Font

# Para checks de tipo Django model
from django.db import models as djmodels

eq_re = re.compile(r'^\s*={1,}\s*(.+?)\s*={1,}\s*$')


def _sanitize_sheet_title(title: str) -> str:
    """
    Sanitiza el título de la hoja (max 31 chars, sin saltos ni barras).
    """
    if not title:
        return "cuestionario"
    title = str(title).replace("\n", " ").replace("/", "-").replace("\\", "-")
    return title[:31]


def _get_field_verbose_name_from_cpqol(cpqol, attr):
    """
    Intenta obtener verbose_name del field 'attr' definido en Cpqol.
    Si no existe, devuelve una versión legible del attr.
    """
    try:
        f = cpqol._meta.get_field(attr)
        return str(f.verbose_name)
    except Exception:
        # fallback: attr con espacio y capitalizado
        return str(attr).replace('_', ' ').capitalize()


def _sanitize_cell_value(val):
    """
    Convertir valores que openpyxl no acepta (ej datetimes con tzinfo) a formatos soportados.
    - datetime con tzinfo -> convertido a UTC naive (tz removed)
    - time con tzinfo -> tzinfo removed
    - Decimal -> float
    - None -> ""
    - bool/int/float/str -> devuelto tal cual
    - otros -> str(...)
    """
    if isinstance(val, datetime):
        if val.tzinfo is not None:
            try:
                # convertir a UTC y quitar tzinfo
                val = val.astimezone(timezone.utc).replace(tzinfo=None)
            except Exception:
                # si falla, solo quitar tzinfo
                val = val.replace(tzinfo=None)
        return val
    if isinstance(val, time):
        # quitar tzinfo si existe
        if getattr(val, "tzinfo", None) is not None:
            try:
                val = val.replace(tzinfo=None)
            except Exception:
                pass
        return val
    if isinstance(val, date) and not isinstance(val, datetime):
        return val
    if isinstance(val, Decimal):
        try:
            return float(val)
        except Exception:
            return str(val)
    if isinstance(val, (bool, int, float, str)):
        return val
    if val is None:
        return ""
    # Si es instancia de modelo, convertir a str
    try:
        if isinstance(val, djmodels.Model):
            return str(val)
    except Exception:
        pass
    # fallback: str()
    try:
        return str(val)
    except Exception:
        return ""


def _format_value(instance, field):
    """
    Devuelve una representación "amigable" del valor y lo sanitiza para openpyxl
    """
    # Caso especial para campo cobertura de Paciente
    if field.name == 'cobertura' and hasattr(instance, 'get_cobertura_display'):
        try:
            val = instance.get_cobertura_display()
            return _sanitize_cell_value(val)
        except Exception:
            pass

    # intentar método get_<field>_display
    getter_name = f'get_{field.name}_display'
    if hasattr(instance, getter_name):
        try:
            val = getattr(instance, getter_name)()
            return _sanitize_cell_value(val)
        except Exception:
            pass

    try:
        val = getattr(instance, field.name)
    except Exception:
        return ""

    # Si es instancia de modelo, usar str()
    try:
        if isinstance(val, djmodels.Model):
            return _sanitize_cell_value(str(val))
    except Exception:
        pass

    # booleano como Sí/No
    if isinstance(val, bool):
        return _sanitize_cell_value("Sí" if val else "No")

    # Manejar listas (como cobertura cuando es JSON)
    if isinstance(val, list):
        try:
            # Si es una lista de opciones, intentar obtener display values
            if hasattr(instance, f'get_{field.name}_display'):
                display_method = getattr(instance, f'get_{field.name}_display')
                return _sanitize_cell_value(display_method())
            else:
                # Convertir lista a string separado por comas
                return _sanitize_cell_value(", ".join(str(item) for item in val))
        except Exception:
            return _sanitize_cell_value(", ".join(str(item) for item in val))

    # por defecto sanitizar
    return _sanitize_cell_value(val)


def workbook_for_cpqols(cpqol_iterable):
    """
    Genera un Workbook que maneja tanto Cpqol como CpqolProfesional
    """
    wb = Workbook()
    # remover la hoja por defecto si existe
    try:
        default = wb.active
        wb.remove(default)
    except Exception:
        pass

    bold = Font(bold=True)

    for cpqol in cpqol_iterable:
        # Determinar el tipo de cuestionario
        is_professional = hasattr(cpqol, 'profesional')
        
        if is_professional:
            # Para profesionales
            sheet_name = _sanitize_sheet_title(getattr(cpqol, "codigo", None) or f"prof_{cpqol.pk}")
            attrs_order = ['profesional', 'contexto', 'datos_clinicos', 'correo']
        else:
            # Para familiares
            sheet_name = _sanitize_sheet_title(getattr(cpqol, "codigo", None) or f"cuest_{cpqol.pk}")
            attrs_order = [
                'tutor', 'paciente', 'movimiento', 'sentimientos', 'relaciones', 
                'familia', 'participacion', 'escuela', 'salud', 'dolor', 'servicios',
                'salud_ultima_semana', 'salud_ultima_semana_2', 'hogar', 'correo'
            ]

        ws = wb.create_sheet(title=sheet_name)

        # ancho de columnas
        ws.column_dimensions['A'].width = 60
        ws.column_dimensions['B'].width = 25

        # Encabezado con el código y tipo del cuestionario
        tipo = "Profesional" if is_professional else "Familiar"
        ws.append([f"Cuestionario {tipo}: {sheet_name}", ""])
        r = ws.max_row
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        ws.cell(row=r, column=1).font = bold
        ws.append([])

        # Campos simples del cuestionario
        try:
            if is_professional:
                simple_fields = ['creacion', 'codigo', 'completado', 'user']
            else:
                simple_fields = ['creacion', 'codigo', 'completado', 'user']
                
            for fname in simple_fields:
                try:
                    f = cpqol._meta.get_field(fname)
                except Exception:
                    continue
                label = str(f.verbose_name)
                value = _format_value(cpqol, f)
                ws.append([label, value])
            ws.append([])
        except Exception:
            pass

        # Recorrer secciones según el tipo
        for attr in attrs_order:
            related = getattr(cpqol, attr, None)
            section_label = _get_field_verbose_name_from_cpqol(cpqol, attr)

            if related is None:
                continue

            # limpiar si el label tiene = ... =
            m = eq_re.match(section_label)
            if m:
                section_label = m.group(1).strip()

            # Escribir título de sección
            ws.append([section_label, ""])
            rr = ws.max_row
            ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=2)
            ws.cell(row=rr, column=1).font = bold

            # Procesar campos del objeto relacionado
            try:
                if hasattr(related, 'all') and not isinstance(related, djmodels.Model):
                    for inst in related.all():
                        inst_label = str(getattr(inst, 'id', inst.pk)) if hasattr(inst, 'pk') else str(inst)
                        ws.append([f"Registro: {inst_label}", ""])
                        rsub = ws.max_row
                        ws.merge_cells(start_row=rsub, start_column=1, end_row=rsub, end_column=2)
                        ws.cell(row=rsub, column=1).font = bold

                        for f in inst._meta.fields:
                            if getattr(f, 'auto_created', False) or isinstance(f, djmodels.AutoField):
                                continue
                            label = str(f.verbose_name)
                            value = _format_value(inst, f)
                            ws.append([label, value])
                        ws.append([])
                else:
                    for f in getattr(related, "_meta", djmodels).fields:
                        if getattr(f, 'auto_created', False) or isinstance(f, djmodels.AutoField):
                            continue
                        label = str(f.verbose_name)
                        value = _format_value(related, f)
                        ws.append([label, value])
            except Exception:
                # fallback para errores
                try:
                    if hasattr(cpqol, 'informacion_completa'):
                        info_text = cpqol.informacion_completa() or ""
                        # ... (código existente para procesar informacion_completa)
                except Exception:
                    ws.append(["(No se pudo obtener detalles de la sección)", ""])

            ws.append([])

    return wb


def workbook_to_bytesio(wb):
    """
    Guarda el workbook en un BytesIO y lo devuelve listo para enviarse.
    """
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio
